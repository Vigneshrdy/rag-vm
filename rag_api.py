
from fastapi import FastAPI, UploadFile, File, HTTPException
from pydantic import BaseModel
from typing import Dict
import os, json
from pathlib import Path
from datetime import datetime, timedelta, timezone

from openai import AzureOpenAI
from supabase import create_client

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_huggingface import HuggingFaceEmbeddings
from fastapi import HTTPException

from pypdf import PdfReader
from docx import Document as DocxDocument

# =====================================================
# ENV (SAFE LOAD)
# =====================================================
AZURE_OPENAI_API_KEY = os.getenv("AZURE_OPENAI_API_KEY")
AZURE_OPENAI_API_VERSION = os.getenv("AZURE_OPENAI_API_VERSION")
AZURE_OPENAI_ENDPOINT = os.getenv("AZURE_OPENAI_ENDPOINT")
AZURE_OPENAI_DEPLOYMENT = os.getenv("AZURE_OPENAI_DEPLOYMENT")

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SERVICE_KEY = os.getenv("SUPABASE_SERVICE_KEY")
APP_URL = os.getenv("APP_URL", "http://localhost:5173")

if not all([
    AZURE_OPENAI_API_KEY,
    AZURE_OPENAI_API_VERSION,
    AZURE_OPENAI_ENDPOINT,
    AZURE_OPENAI_DEPLOYMENT,
    SUPABASE_URL,
    SUPABASE_SERVICE_KEY,
]):
    raise RuntimeError("Missing environment variables")

client = AzureOpenAI(
    api_key=AZURE_OPENAI_API_KEY,
    api_version=AZURE_OPENAI_API_VERSION,
    azure_endpoint=AZURE_OPENAI_ENDPOINT,
)

supabase = create_client(SUPABASE_URL, SUPABASE_SERVICE_KEY)

app = FastAPI(title="Nyaya AI – Full RAG API")

# =====================================================
# STORAGE
# =====================================================
DOCUMENT_INDEX: Dict[str, FAISS] = {}

BASE_DIR = Path(".")
UPLOAD_DIR = BASE_DIR / "uploads"
CHAT_DIR = BASE_DIR / "chats"

UPLOAD_DIR.mkdir(exist_ok=True)
CHAT_DIR.mkdir(exist_ok=True)

# =====================================================
# SHARE CHAT ENDPOINTS
# =====================================================

@app.post("/api/share-chat")
def share_chat(payload: dict):
    session_id = payload.get("session_id")
    if not session_id:
        raise HTTPException(400, "session_id is required")

    # 1️⃣ Try to find existing share
    existing = (
        supabase.table("shared_chats")
        .select("id")
        .eq("session_id", session_id)
        .limit(1)
        .execute()
    )

    # ✅ SAFE CHECK
    if existing and existing.data:
        share_id = existing.data[0]["id"]
    else:
        # 2️⃣ Fetch messages for this session
        chat = (
            supabase.table("chats")
            .select("question, answer, source, created_at")
            .eq("session_id", session_id)
            .order("created_at")
            .execute()
        )

        if not chat.data:
            raise HTTPException(404, "No chat found for this session")

        expires_at = datetime.now(timezone.utc) + timedelta(days=7)

        insert = (
            supabase.table("shared_chats")
            .insert({
                "session_id": session_id,
                "messages": chat.data,  # full chat history
                "expires_at": expires_at.isoformat(),
                "is_public": True,
            })
            .execute()
        )

        share_id = insert.data[0]["id"]

    # 3️⃣ ALWAYS return stable URL
    return {
        "share_url": f"https://nyayaai.saireddy.dev/share/{share_id}"
    }

@app.get("/api/share/{share_id}")
@app.get("/api/share/{share_id}")
def get_shared_chat(share_id: str):
    result = (
        supabase.table("shared_chats")
        .select("messages, expires_at")
        .eq("id", share_id)
        .single()
        .execute()
    )

    if not result.data:
        raise HTTPException(404, "Chat not found")

    expires_at = result.data["expires_at"]

    if expires_at:
        expires = datetime.fromisoformat(expires_at)

        # 🔥 FORCE timezone awareness
        if expires.tzinfo is None:
            expires = expires.replace(tzinfo=timezone.utc)

        if datetime.now(timezone.utc) > expires:
            raise HTTPException(410, "Chat expired")

    return {"messages": result.data["messages"]}

# =====================================================
# EMBEDDINGS
# =====================================================
embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")

# =====================================================
# UTILS
# =====================================================

def extract_text(file, filename: str) -> str:
    if filename.lower().endswith(".pdf"):
        reader = PdfReader(file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if filename.lower().endswith(".docx"):
        doc = DocxDocument(file)
        return "\n".join(p.text for p in doc.paragraphs)
    if filename.lower().endswith(".txt"):
        return file.read().decode("utf-8")
    raise ValueError("Unsupported file type")



# =====================================================
# MAIN API CONTINUES...
# (Your /api/upload, /api/ask, /api/ask-document)
# unchanged
# =====================================================



def save_chat_local(session_id: str, role: str, content: str):
    path = CHAT_DIR / f"{session_id}.json"
    data = json.loads(path.read_text()) if path.exists() else []
    data.append({"role": role, "content": content})
    path.write_text(json.dumps(data, indent=2))


def save_chat_db(session_id: str, source: str, question: str, answer: str):
    supabase.table("chats").insert({
        "session_id": session_id,
        "source": source,
        "question": question,
        "answer": answer
    }).execute()


def is_legal_document(text: str) -> dict:
    resp = client.chat.completions.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        messages=[
            {
                "role": "system",
                "content": (
                    "Classify if the document is a legal document under Indian law.\n"
                    "Legal documents include judgments, FIRs, contracts, petitions, notices, Acts.\n"
                    "Reply STRICTLY in JSON:\n"
                    "{ \"legal\": true/false, \"reason\": \"short reason\" }"
                )
            },
            {"role": "user", "content": text[:3000]}
        ]
    )
    return json.loads(resp.choices[0].message.content)

# =====================================================
# UPLOAD + SUMMARY (RAG INGESTION)
# =====================================================
@app.post("/api/upload")
async def upload(session_id: str, file: UploadFile = File(...)):
    # Save file
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(exist_ok=True)

    file_path = session_dir / file.filename
    with open(file_path, "wb") as f:
        f.write(await file.read())

    with open(file_path, "rb") as f:
        text = extract_text(f, file.filename)

    # Legal relevance check
    verdict = is_legal_document(text)

    # Insert metadata first
    insert_resp = supabase.table("documents").insert({
        "session_id": session_id,
        "filename": file.filename,
        "file_path": str(file_path),
        "is_legal": verdict["legal"],
        "reason": verdict["reason"]
    }).execute()

    document_id = insert_resp.data[0]["id"]

    if not verdict["legal"]:
        raise HTTPException(400, f"Rejected: {verdict['reason']}")

    # Chunk + embed
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=150
    )

    docs = splitter.split_documents([Document(page_content=text)])
    vectorstore = FAISS.from_documents(docs, embeddings)
    DOCUMENT_INDEX[session_id] = vectorstore

    # Retrieve best chunks for summary
    context = "\n\n".join(
        d.page_content
        for d in vectorstore.similarity_search("summarize judgment", k=6)
    )

    # Generate summary
    resp = client.chat.completions.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        messages=[
            {
                "role": "system",
                "content": (
                    "Summarize the legal document in Markdown.\n"
                    "Format strictly as:\n"
                    "### Facts\n"
                    "### Issues\n"
                    "### Decision\n"
                    "### Relevant Provisions"
                )
            },
            {"role": "user", "content": context}
        ]
    )

    summary = resp.choices[0].message.content

    # Save summary to documents table
    supabase.table("documents").update({
        "summary": summary
    }).eq("id", document_id).execute()

    save_chat_local(session_id, "assistant", summary)

    return {"summary": summary}

# =====================================================
# ASK DOCUMENT (RAG QA)
# =====================================================
class DocQuery(BaseModel):
    session_id: str
    question: str

@app.post("/api/ask-document")
def ask_document(q: DocQuery):
    if q.session_id not in DOCUMENT_INDEX:
        return {"answer": "No document uploaded for this session."}

    save_chat_local(q.session_id, "user", q.question)

    docs = DOCUMENT_INDEX[q.session_id].similarity_search(q.question, k=5)
    context = "\n\n".join(d.page_content for d in docs)

    resp = client.chat.completions.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        messages=[
            {
                "role": "system",
                "content": (
                    "Answer ONLY from the document.\n"
                    "Use Markdown formatting.\n"
                    "If not present, reply exactly: Not mentioned in the document."
                )
            },
            {"role": "user", "content": context + "\n\nQ: " + q.question}
        ]
    )

    answer = resp.choices[0].message.content

    save_chat_local(q.session_id, "assistant", answer)
    save_chat_db(q.session_id, "document", q.question, answer)

    return {"answer": answer}

# =====================================================
# NORMAL CHAT
# =====================================================
class AskQuery(BaseModel):
    session_id: str
    question: str

@app.post("/api/ask")
def ask(q: AskQuery):
    save_chat_local(q.session_id, "user", q.question)

    resp = client.chat.completions.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are Nyaya AI, an Indian legal assistant.\n"
                    "Respond in Markdown.\n"
                    "Do not give legal advice."
                )
            },
            {"role": "user", "content": q.question}
        ]
    )

    answer = resp.choices[0].message.content

    save_chat_local(q.session_id, "assistant", answer)
    save_chat_db(q.session_id, "normal", q.question, answer)

    return {"answer": answer}
